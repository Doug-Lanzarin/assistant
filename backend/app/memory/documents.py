"""Biblioteca pessoal: importação e extração de texto de documentos.

Formatos: txt, md, código-fonte, csv, json, html, pdf (pypdf), docx (python-docx).
O conteúdo importado é tratado como DADO NÃO CONFIÁVEL: nunca vira instrução.
"""

from __future__ import annotations

import hashlib
import io
import re
from pathlib import Path

from app.core.db import Database, from_json, to_json, utcnow
from app.memory.graph import Graph
from app.memory.index import Indexer

TEXT_EXT = {
    ".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json", ".yaml", ".yml", ".toml", ".ini", ".cfg",
    ".log", ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".kt", ".cs", ".go", ".rs", ".rb", ".php", ".sql",
    ".sh", ".ps1", ".bat", ".c", ".h", ".cpp", ".hpp", ".css", ".scss", ".xml", ".vue", ".swift", ".dart",
}
MAX_BYTES = 25 * 1024 * 1024


class UnsupportedDocument(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext in TEXT_EXT:
        for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
            try:
                return data.decode(enc)
            except UnicodeDecodeError:
                continue
        raise UnsupportedDocument("Codificação de texto não reconhecida")
    if ext in {".html", ".htm"}:
        text = data.decode("utf-8", errors="replace")
        text = re.sub(r"(?is)<(script|style).*?</\1>", " ", text)
        text = re.sub(r"(?s)<[^>]+>", " ", text)
        return re.sub(r"[ \t]+", " ", text)
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = []
        for i, page in enumerate(reader.pages):
            try:
                pages.append(f"[página {i + 1}]\n" + (page.extract_text() or ""))
            except Exception:  # noqa: BLE001 - página corrompida não derruba a importação
                pages.append(f"[página {i + 1}: falha na extração]")
        return "\n\n".join(pages)
    if ext == ".docx":
        import docx

        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    raise UnsupportedDocument(f"Formato não suportado: {ext or 'sem extensão'}")


class Library:
    def __init__(self, db: Database, indexer: Indexer, graph: Graph, storage_dir: Path):
        self.db = db
        self.indexer = indexer
        self.graph = graph
        self.storage = storage_dir

    async def import_bytes(
        self, filename: str, data: bytes, title: str | None = None, project: str | None = None,
        tags: list[str] | None = None, project_id: int | None = None,
    ) -> dict:
        if len(data) > MAX_BYTES:
            raise UnsupportedDocument("Arquivo maior que 25 MB")
        safe_name = Path(filename).name  # impede path traversal
        text = extract_text(safe_name, data).strip()
        if not re.sub(r"\[página \d+(: falha na extração)?\]", "", text).strip():
            raise UnsupportedDocument("Nenhum texto extraído (PDF escaneado? use OCR antes)")
        sha = hashlib.sha256(data).hexdigest()
        existing = self.db.query_one("SELECT id FROM documents WHERE sha256=?", (sha,))
        if existing:
            return {**self.get(existing["id"]), "duplicate": True}
        self.storage.mkdir(parents=True, exist_ok=True)
        stored = self.storage / f"{sha[:16]}_{safe_name}"
        stored.write_bytes(data)
        now = utcnow()
        did = self.db.insert(
            "INSERT INTO documents(title, filename, stored_path, mime, sha256, size, project_id, tags, char_count,"
            " imported_at, updated_at) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
            (title or Path(safe_name).stem, safe_name, str(stored), Path(safe_name).suffix.lower(), sha, len(data),
             project_id, to_json(tags or []), len(text), now, now),
        )
        chunks = await self.indexer.index("document", did, text)
        proj_name = None
        if project_id:
            row = self.db.query_one("SELECT name FROM projects WHERE id=?", (project_id,))
            proj_name = row["name"] if row else None
        self.graph.index_owner("document", did, text[:20000], "document", proj_name, title or safe_name)
        return {**self.get(did), "chunks": chunks, "duplicate": False}

    async def import_path(self, path: Path, **kw) -> dict:
        return await self.import_bytes(path.name, path.read_bytes(), **kw)

    def get(self, did: int) -> dict:
        d = self.db.query_one(
            "SELECT d.*, p.name AS project_name FROM documents d LEFT JOIN projects p ON p.id=d.project_id WHERE d.id=?",
            (did,),
        )
        if not d:
            raise KeyError("Documento não encontrado")
        d["tags"] = from_json(d["tags"], [])
        d["chunk_count"] = self.db.scalar("SELECT COUNT(*) FROM chunks WHERE owner_type='document' AND owner_id=?", (did,))
        return d

    def list(self) -> list[dict]:
        rows = self.db.query(
            "SELECT d.id, d.title, d.filename, d.size, d.char_count, d.imported_at, p.name AS project_name"
            " FROM documents d LEFT JOIN projects p ON p.id=d.project_id ORDER BY d.imported_at DESC"
        )
        return rows

    def text(self, did: int, max_chars: int = 50_000) -> str:
        rows = self.db.query("SELECT text FROM chunks WHERE owner_type='document' AND owner_id=? ORDER BY ord", (did,))
        return "\n".join(r["text"] for r in rows)[:max_chars]

    def delete(self, did: int) -> dict[str, int]:
        d = self.get(did)
        with self.db.transaction():
            chunks = self.indexer.remove("document", did)
            links = self.graph.unlink_owner("document", did)
            self.db.execute("DELETE FROM documents WHERE id=?", (did,))
        if d.get("stored_path"):
            p = Path(d["stored_path"])
            if p.exists() and self.storage in p.resolve().parents:
                p.unlink()
        return {"documents": 1, "chunks": chunks, "graph_links": links}

